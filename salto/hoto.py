"""HOTO preparation: the hand-over sheet of a LAN project, one row per device in creation."""
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .flip import norm

COLUMNS = ["Gold Ref", "Go Live Date", "Project Manager", "City", "Country", "Site address", "Site ID",
           "Device Name", "OBS Device name", "Device Type", "Switch number", "Stack Master Name", "Device Model",
           "USID", "Admin ip address", "Customer ip address", "Device SN#", "Local site contact",
           "Telephone / mobile number of local site person", "Email Address of local contact"]

# SALTO device type -> HOTO short type
TYPES = {"lan switch": "SW", "switch": "SW", "wireless access point": "AP", "access point": "AP",
         "wireless lan controller": "WLC", "router": "RTR", "firewall": "FW"}
# device name code (skf-<code>-site-nn) -> HOTO short type
CODES = {"sw": "SW", "ap": "AP", "wlc": "WLC", "rt": "RTR", "rtr": "RTR", "fw": "FW"}


def is_lan(rep):
    case = rep.case
    product = dict(case.header).get("Product", "") if case else ""
    return bool(case) and (rep.stc_id.startswith("MLAN-") or "lan" in product.lower())


def _fields(case, tab, section=None):
    t = case.tabs.get(tab)
    if not t:
        return {}
    out = {}
    for name, pairs in t.sections:
        if section is None or name.lower() == section.lower():
            for k, v in pairs:
                out.setdefault(k, v)
    return out


def _date(text):
    for fmt in ("%d-%b-%Y", "%d %b %Y", "%d/%m/%Y"):
        try:
            d = datetime.strptime(text.strip(), fmt)
            return f"{d.month}/{d.day}/{d.year}"
        except ValueError:
            pass
    return text


def creation_devices(case):
    """Scope-of-work devices being created (not ceased or modified)."""
    devs = case.scope_devices
    if any(d.get("Cycle") for d in devs):
        return [d for d in devs if (d.get("Cycle") or "").lower() == "creation"]
    return [d for d in devs if (d.get("Cycle") or "").lower() != "cessation"]


def device_type(dev):
    t = TYPES.get((dev.get("Type") or "").strip().lower())
    if t:
        return t
    parts = (dev.get("Customer Defined Code") or "").lower().split("-")
    return CODES.get(parts[1], parts[1].upper()) if len(parts) > 2 else dev.get("Type", "")


def device_number(name):
    """skf-sw-itvil-01_01 -> 01_01, skf-sw-itvil-02_01-a14 -> 02_01, skf-ap-semol-03 -> 03."""
    name = (name or "").strip()
    m = re.search(r"-(\d+_\d+)(?=-|$)", name) or re.search(r"-(\d+)(?=-|$)", name)
    return m.group(1) if m else ""


def build(rep):
    """HOTO rows for the order's devices in creation (Admin / Customer IP are not in SALTO: left empty)."""
    case = rep.case
    info = _fields(case, "Information", "MAIN DATES")
    go_live = info.get("Revised Target Delivery Date") or info.get("Initial Target Delivery Date") or \
        dict(case.header).get("Target Delivery Date", "")
    pm = next((p.name for p in case.participants if p.role.lower() == "project manager"), "")
    site, addr = _fields(case, "Site", "SITE"), _fields(case, "Site", "Address")
    contact = _fields(case, "Site", "Primary Contact")
    person = " ".join(x for x in (contact.get("First Name", ""), contact.get("Last Name", "")) if x)
    phone = contact.get("Mobile") or contact.get("Phone", "")
    serials = {norm(s["Device (OBS name)"]): s["Serial number"]
               for c in case.walk() for s in (c.flip or {}).get("serials", []) if s["Serial number"]}

    devs = creation_devices(case)
    rows = []
    for dev in devs:
        name = dev.get("Customer Defined Code", "")
        parts = name.split("-")
        site_id = (site.get("Customer Site Name") or (parts[2] if len(parts) > 3 else "")).upper()
        kind = device_type(dev)
        number = device_number(name) if kind == "SW" else ""
        master = ""
        if "_" in number:  # stack: <stack>_<member>, member 01 is the master
            stack, member = number.split("_", 1)
            if int(member) == 1:
                master = "Master"
            else:
                head = name.replace(f"-{number}", f"-{stack}_{'1'.zfill(len(member))}", 1)
                master = head if any(d.get("Customer Defined Code") == head for d in devs) else "Member"
        rows.append({
            "Gold Ref": rep.order_ref, "Go Live Date": _date(go_live), "Project Manager": pm,
            "City": addr.get("City", ""), "Country": addr.get("Country", "").upper(),
            "Site address": addr.get("Address", ""), "Site ID": site_id,
            "Device Name": name, "OBS Device name": dev.get("Name", ""), "Device Type": kind,
            "Switch number": number, "Stack Master Name": master, "Device Model": dev.get("Model", ""),
            "USID": dev.get("USID", ""), "Admin ip address": "", "Customer ip address": "",
            "Device SN#": serials.get(norm(dev.get("Name", "")), ""),
            "Local site contact": person, "Telephone / mobile number of local site person": phone,
            "Email Address of local contact": contact.get("Email", ""),
        })
    return sorted(rows, key=lambda r: (r["Device Type"] != "SW", r["Device Name"]))


def needs_check(row):
    return not row["Device SN#"] or " / " in row["Device SN#"]


DARK, ORANGE, TINT, BAND, INPUT, CHECK, LINE = "2E3B3B", "FF7900", "FFE8D6", "F7F8FA", "FFFBEA", "FFE7A3", "D9DDE3"
TYPE_NAMES = {"SW": "Switches", "AP": "Access points", "WLC": "Wireless LAN controllers", "RTR": "Routers",
              "FW": "Firewalls"}
INPUT_COLUMNS = ("Admin ip address", "Customer ip address")  # not in SALTO: filled in by hand


def save(rows, path, rep=None):
    """The HOTO workbook: a designed 'HOTO' sheet (title, order block, devices grouped by type, totals) and a
    'Plain list' sheet with the same rows, ready to paste into another HOTO document."""
    from openpyxl.styles import Border, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=LINE)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_bold = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    n = len(COLUMNS)
    last = get_column_letter(n)
    first = rows[0] if rows else {}
    header = dict(rep.case.header) if rep is not None and getattr(rep, "case", None) else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "HOTO"
    ws.sheet_view.showGridLines = False
    # title band
    ws.merge_cells(f"A1:{get_column_letter(n - 3)}1")
    ws["A1"] = "HOTO - Handover to Operations"
    ws["A1"].font = Font(name="Calibri", bold=True, size=20, color="1F2328")
    ws.merge_cells(f"{get_column_letter(n - 2)}1:{last}1")
    ws[f"{get_column_letter(n - 2)}1"] = "Prepared with SALTO Buddy"
    ws[f"{get_column_letter(n - 2)}1"].font = Font(italic=True, color="6B7280", size=9)
    ws[f"{get_column_letter(n - 2)}1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{last}2")
    ws["A2"].fill = PatternFill("solid", fgColor=ORANGE)
    ws.row_dimensions[2].height = 4
    # order block: labels on a dark band, values in boxes below
    info = [("GOLD REF", first.get("Gold Ref", "")), ("CUSTOMER", header.get("Customer name", "")),
            ("SITE", f"{first.get('Site ID', '')} - {first.get('City', '')}, {first.get('Country', '')}".strip(" -,")),
            ("GO LIVE DATE", first.get("Go Live Date", "")), ("PROJECT MANAGER", first.get("Project Manager", "")),
            ("STC CASE", getattr(rep, "stc_id", "") if rep is not None else ""),
            ("DATE OF LAST UPDATE", datetime.now().strftime("%m/%d/%Y"))]
    span = max(1, n // len(info))
    for i, (label, value) in enumerate(info):
        c1 = 1 + i * span
        c2 = n if i == len(info) - 1 else c1 + span - 1
        for row, val, style in ((4, label, "label"), (5, value, "value")):
            if c2 > c1:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            c = ws.cell(row=row, column=c1, value=val)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if style == "label":
                c.font = Font(bold=True, color="FFFFFF", size=9)
            else:
                c.font = Font(bold=True, size=11, color="1F2328")
            for col in range(c1, c2 + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill("solid", fgColor=DARK if style == "label" else "FFFFFF")
                cell.border = box
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 26
    ws.cell(row=7, column=1, value=f"{len(rows)} device(s) in creation. Yellow columns are filled in by hand; "
                                   "amber serial numbers need checking (missing, or several in FLIP).").font =         Font(italic=True, color="6B7280", size=9)

    # devices table
    top = 9
    for j, h in enumerate(COLUMNS, 1):
        c = ws.cell(row=top, column=j, value=h)
        c.fill = PatternFill("solid", fgColor=DARK)
        c.font = white_bold
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[top].height = 42
    r = top + 1
    groups = {}
    for row in rows:
        groups.setdefault(row.get("Device Type", ""), []).append(row)
    for kind, items in groups.items():
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        c = ws.cell(row=r, column=1, value=f"{TYPE_NAMES.get(kind, kind or 'Devices').upper()}  ({len(items)})")
        c.font = Font(bold=True, size=10, color="1F2328")
        for col in range(1, n + 1):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=TINT)
            ws.cell(row=r, column=col).border = box
        ws.row_dimensions[r].height = 20
        r += 1
        for k, row in enumerate(items):
            for j, h in enumerate(COLUMNS, 1):
                c = ws.cell(row=r, column=j, value=row.get(h, ""))
                c.border = box
                c.font = Font(size=10)
                c.alignment = Alignment(vertical="center", wrap_text=h in ("Site address",))
                if h in INPUT_COLUMNS:
                    c.fill = PatternFill("solid", fgColor=INPUT)
                elif h == "Device SN#" and needs_check(row):
                    c.fill = PatternFill("solid", fgColor=CHECK)
                elif k % 2:
                    c.fill = PatternFill("solid", fgColor=BAND)
            r += 1
    # totals
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n - 2)
    ws.cell(row=r, column=1, value="TOTAL DEVICES IN CREATION")
    ws.merge_cells(start_row=r, start_column=n - 1, end_row=r, end_column=n)
    ws.cell(row=r, column=n - 1, value=len(rows))
    for col in range(1, n + 1):
        c = ws.cell(row=r, column=col)
        c.fill = PatternFill("solid", fgColor=DARK)
        c.font = white_bold
        c.border = box
        c.alignment = Alignment(horizontal="right" if col >= n - 1 else "left", vertical="center")
    ws.row_dimensions[r].height = 22

    for j, h in enumerate(COLUMNS, 1):
        width = max([min(len(h), 18)] + [len(str(x.get(h, ""))) for x in rows])
        ws.column_dimensions[get_column_letter(j)].width = min(max(11, width + 2), 40)
    ws.freeze_panes = f"C{top + 1}"  # keep the header row and the first two columns in view
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # the same rows, unstyled, for pasting into another HOTO document
    plain = wb.create_sheet("Plain list")
    for j, h in enumerate(COLUMNS, 1):
        plain.cell(row=1, column=j, value=h).font = Font(bold=True)
        plain.column_dimensions[get_column_letter(j)].width = ws.column_dimensions[get_column_letter(j)].width
    for i, row in enumerate(rows, 2):
        for j, h in enumerate(COLUMNS, 1):
            plain.cell(row=i, column=j, value=row.get(h, ""))
    plain.freeze_panes = "A2"
    wb.save(path)
    return path
