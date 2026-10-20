"""Export an OrderReport to Excel (one sheet per area)."""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="FF7900")  # Orange
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=12)


def _sheet(wb, title):
    title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31] or "Sheet"
    base, n = title, 2
    while title in wb.sheetnames:
        title = f"{base[:28]}~{n}"
        n += 1
    return wb.create_sheet(title)


def _title(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    return row + 1


def _table(ws, row, headers, rows):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        row += 1
        for j, h in enumerate(headers, 1):
            ws.cell(row=row, column=j, value=r.get(h, "") if isinstance(r, dict) else r[j - 1])
    return row + 2


def _autosize(ws):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), min(60, len(str(c.value))))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w + 2
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)


def export_report(rep, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = _title(ws, 1, f"Order {rep.order_ref}")
    summary = [("STC ID", rep.stc_id), ("Status", rep.status), ("Current stage", rep.current_stage),
               ("Stages", " > ".join(f"{s.name} ({s.state})" for s in rep.stages)),
               ("EQPs", f"{len(rep.eqps)}: {', '.join(rep.eqps)}")]
    case = rep.case
    if case:
        summary.append(("Scope of work devices", str(len(case.scope_devices))))
        summary += case.header
    row = _table(ws, row, ["Field", "Value"], summary)
    if case:
        row = _title(ws, row, "Participants")
        row = _table(ws, row, ["Name", "Role", "Party"], [(p.name, p.role, p.party) for p in case.participants])
    _autosize(ws)

    if case:
        ws = _sheet(wb, "Scope of work")
        row = 1
        for n in case.scope:
            row = _title(ws, row, f"{n.product} > {n.node}")
            for k, v in n.counts:
                ws.cell(row=row, column=1, value=k)
                ws.cell(row=row, column=2, value=v)
                row += 1
            row = _table(ws, row + 1, n.devices.headers, n.devices.rows)
        _autosize(ws)
        for name, td in case.tabs.items():
            if name == "TS":
                continue
            ws = _sheet(wb, name)
            row = 1
            for sec, fields in td.sections:
                row = _title(ws, row, sec)
                row = _table(ws, row, ["Field", "Value"], fields)
            for g in td.grids:
                row = _title(ws, row, g.name)
                row = _table(ws, row, g.headers, g.rows)
            if row == 1:
                ws.cell(row=1, column=1, value="No data")
            _autosize(ws)

    ws = _sheet(wb, "Search results")
    width = max((len(r.cells) for r in rep.rows), default=0)
    _table(ws, 1, ["STC ID", "Sub-case ID"] + [f"Col {i + 1}" for i in range(width)],
           [[r.stc_id, r.sub_case_id] + r.cells + [""] * (width - len(r.cells)) for r in rep.rows])
    _autosize(ws)
    wb.save(path)
    return path
